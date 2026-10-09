"""전 종목 스크리너(src/screener.py · src/screen_excel.py · screen.py) 테스트 — 네트워크 없음."""

from __future__ import annotations

import datetime as dt

import pytest
from openpyxl import load_workbook

import screen
from src import screener
from src.models import SubCheck
from src.screen_excel import write_excel
from src.screener import NEAR, PASS, UniverseItem
from tests.conftest import (
    SAMSUNG_A_EPS, SAMSUNG_A_PERIODS, SAMSUNG_A_REVENUE, SAMSUNG_A_ROE,
    SAMSUNG_Q_EPS, SAMSUNG_Q_PERIODS, SAMSUNG_Q_REVENUE, SAMSUNG_Q_ROE, make_finance,
)

ITEM = UniverseItem("005930", "삼성전자", "KOSPI", 15_375_713.0 * 1e8)


def stock(code, name, end="stock", cap="1,000"):
    return {"itemCode": code, "stockName": name, "stockEndType": end, "marketValue": cap}


def finance_payloads(q_roe=None, q_revenue=None):
    q = make_finance(SAMSUNG_Q_PERIODS, {"매출액": q_revenue or SAMSUNG_Q_REVENUE,
                                         "EPS": SAMSUNG_Q_EPS, "ROE": q_roe or SAMSUNG_Q_ROE})
    a = make_finance(SAMSUNG_A_PERIODS, {"매출액": SAMSUNG_A_REVENUE, "EPS": SAMSUNG_A_EPS,
                                         "ROE": SAMSUNG_A_ROE})
    return {"quarter": q, "annual": a}


@pytest.fixture
def fake_finance(monkeypatch):
    def install(**kw):
        payloads = finance_payloads(**kw)
        monkeypatch.setattr(screener.naver, "finance", lambda code, period: payloads[period])
    return install


class Test대상종목:
    def test_ETF_우선주_스팩은_뺀다(self):
        payload = {"stocks": [
            stock("005930", "삼성전자"), stock("069500", "KODEX 200", end="etf"),
            stock("005935", "삼성전자우"), stock("123450", "엔에이치스팩34호"),
            stock("000660", "SK하이닉스", cap="12,316,101"),
        ]}
        items = screener.parse_universe(payload, "KOSPI")
        assert [i.name for i in items] == ["삼성전자", "SK하이닉스"]
        assert items[1].cap == 12_316_101.0 * 1e8   # 억원 → 원
        assert items[0].ref.market == "코스피" and items[0].ref.region == "KR"

    def test_이름에_우가_들어가도_보통주면_남긴다(self):
        items = screener.parse_universe({"stocks": [stock("093240", "에코글로우")]}, "KOSDAQ")
        assert len(items) == 1

    def test_페이지를_끝까지_받는다(self, monkeypatch):
        pages = {1: [stock("000010", "가")], 2: [stock("000020", "나")], 3: []}
        calls = []

        def market_list(market, page, size):
            calls.append((market, page))
            return {"stocks": pages[page]}
        monkeypatch.setattr(screener.naver, "market_list", market_list)
        items = screener.fetch_universe(("KOSDAQ",))
        assert [i.name for i in items] == ["가", "나"]
        assert calls == [("KOSDAQ", 1), ("KOSDAQ", 2), ("KOSDAQ", 3)]


class Test상태분류:
    def check(self, code, passed):
        return SubCheck(code, code, "", passed)

    def test_C2_판단불가는_통과를_막지_않는다(self):
        status, failed, unknown = screener.classify(
            [self.check("C1", True), self.check("C2", None), self.check("A1", True)])
        assert status == PASS and unknown == []

    def test_다른_조건의_판단불가는_근접(self):
        status, _, unknown = screener.classify([self.check("C1", True), self.check("A1", None)])
        assert status == NEAR and unknown == ["A1"]

    def test_불합격_1개는_근접_2개는_탈락(self):
        assert screener.classify([self.check("C4", False), self.check("A1", True)])[0] == NEAR
        assert screener.classify([self.check("C4", False), self.check("A3", False)])[0] == screener.FAIL


class Test1차판정:
    def test_삼성_fixture_는_통과(self, fake_finance):
        fake_finance()
        r = screener.screen_ca(ITEM)
        assert r.status == PASS, r.detail
        assert r.latest_quarter == "2026.03"
        assert round(r.eps_yoy, 1) == round((6993 / 1186 - 1) * 100, 1)
        assert round(r.revenue_yoy, 1) == round((1_338_734 / 791_405 - 1) * 100, 1)
        assert r.roe == 19.16
        assert "C1" in r.detail and "A3" in r.detail

    def test_ROE_미달_하나면_근접(self, fake_finance):
        fake_finance(q_roe={**SAMSUNG_Q_ROE, "202603": "10.00"})
        r = screener.screen_ca(ITEM)
        assert r.status == NEAR and r.failed == ["A3"]

    def test_매출도_정체면_탈락(self, fake_finance):
        fake_finance(q_roe={**SAMSUNG_Q_ROE, "202603": "10.00"},
                     q_revenue={**SAMSUNG_Q_REVENUE, "202603": "800,000"})
        r = screener.screen_ca(ITEM)
        assert r.status == screener.FAIL and set(r.failed) == {"C4", "A3"}

    def test_조회_실패는_오류로_기록하고_계속(self, monkeypatch):
        def boom(code, period):
            raise RuntimeError("HTTP 500")
        monkeypatch.setattr(screener.naver, "finance", boom)
        r = screener.screen_ca(ITEM)
        assert r.status == screener.ERROR and "500" in r.error


class Test실행:
    def universe(self):
        return [UniverseItem(f"{i:05d}0", f"종목{i}", "KOSPI", cap) for i, cap in enumerate([900, 500, 100])]

    def test_시총_하한과_개수_제한(self, fake_finance):
        fake_finance()
        r = screener.run(universe=self.universe(), min_cap=200, limit=1, deep=False, progress=lambda m: None)
        assert r.universe == 1 and [x.item.name for x in r.ca] == ["종목0"]

    def test_2차는_통과_종목만(self, fake_finance, monkeypatch):
        fake_finance()
        seen = []
        monkeypatch.setattr(screener, "deep_check",
                            lambda item: seen.append(item.code) or screener.DeepResult(item, summary="미충족"))
        r = screener.run(universe=self.universe(), deep=True, progress=lambda m: None)
        assert len(seen) == len(r.by_status(PASS)) == 3

    def test_중단하면_그때까지_결과를_돌려준다(self, monkeypatch):
        calls = {"n": 0}

        def screen_ca(item):
            calls["n"] += 1
            if calls["n"] == 2:
                raise KeyboardInterrupt
            return screener.CaResult(item, PASS)
        monkeypatch.setattr(screener, "screen_ca", screen_ca)
        r = screener.run(universe=self.universe(), deep=False, progress=lambda m: None)
        assert r.interrupted and len(r.ca) == 1


class Test엑셀:
    def test_시트와_행(self, tmp_path, fake_finance, monkeypatch):
        fake_finance()
        items = [ITEM, UniverseItem("000660", "SK하이닉스", "KOSPI", 1.0)]
        verdicts = dict.fromkeys("CANSLIM", "합격")
        monkeypatch.setattr(screener, "deep_check", lambda item: screener.DeepResult(
            item, summary="CANSLIM 충족", tally="합격 7", qualified=item.code == "005930",
            verdicts=verdicts, opinion="🟢 매수", base_gap=12.3))
        result = screener.run(universe=items, deep=True, progress=lambda m: None)
        path = write_excel(result, tmp_path / "out.xlsx")

        wb = load_workbook(path)
        assert wb.sheetnames == ["요약", "CANSLIM 충족", "C·A 통과", "C·A 근접"]
        qualified = list(wb["CANSLIM 충족"].iter_rows(min_row=2, values_only=True))
        assert [row[1] for row in qualified] == ["삼성전자"]
        passed = wb["C·A 통과"]
        assert passed.max_row == 3                     # 머리글 + 2종목
        head = [c.value for c in passed[1]]
        c_col = head.index("C") + 1
        assert passed.cell(row=2, column=c_col).fill.fgColor.rgb.endswith("D7F2DF")  # 합격 초록
        summary = {r[0]: r[1] for r in wb["요약"].iter_rows(values_only=True)}
        assert summary["2차 CANSLIM 충족"].startswith("1개")
        assert "+17%" in summary["1차 기준 (웹앱과 같음)"]


class Test실행파일:
    def test_같은_날_파일은_덮어쓰지_않는다(self, tmp_path, monkeypatch):
        monkeypatch.setattr(screen, "OUT_DIR", tmp_path)
        day = dt.date(2026, 10, 9)
        first = screen.output_path(day)
        first.write_bytes(b"x")
        assert screen.output_path(day).name == "CANSLIM_스크리닝_2026-10-09_2.xlsx"

    def test_옵션(self):
        args = screen.parse_args(["--market", "KOSDAQ", "--limit", "5", "--min-cap", "1000", "--no-deep"])
        assert (args.market, args.limit, args.min_cap, args.no_deep) == ("KOSDAQ", 5, 1000.0, True)


# ------------------------------------------------------------------ 미국

def us_row(symbol, name, cap="5000000000"):
    return {"symbol": symbol, "name": name, "marketCap": cap}


class Test미국대상:
    def test_보통주와_ADR만_시총순(self):
        payload = {"data": {"rows": [
            us_row("AAPL", "Apple Inc. Common Stock", "3000000000000"),
            us_row("TSM", "Taiwan Semiconductor Manufacturing Company Ltd. American Depositary Shares",
                   "900000000000"),
            us_row("BRK/B", "Berkshire Hathaway Inc. Class B", "1000000000000"),
            us_row("ABCDW", "Abc Corp Warrant"),
            us_row("XYZ^A", "Xyz Corp Preferred Series A"),
            us_row("UNIT", "Some Acquisition Units"),
            us_row("ZERO", "Zero Cap Inc.", ""),
            us_row("U", "Unity Software Inc. Common Stock"),
        ]}}
        items = screener.parse_universe_us(payload)
        assert [i.code for i in items] == ["AAPL", "BRK-B", "TSM", "U"]
        assert items[0].region == "US" and items[0].currency == "USD"
        assert items[0].ref.region == "US" and items[0].cap == 3e12

    def test_시총_하한은_억달러_단위(self):
        assert screen.min_cap_base("US", None) == 20e8        # 기본 20억 달러
        assert screen.min_cap_base("US", 100) == 100e8
        assert screen.min_cap_base("ALL", None) == 0.0
        assert screen.min_cap_base("KOSPI", 1000) == 1000e8   # 1,000억원

    def test_결과_파일_이름(self, tmp_path, monkeypatch):
        monkeypatch.setattr(screen, "OUT_DIR", tmp_path)
        assert screen.output_path(dt.date(2026, 10, 9), us=True).name == "CANSLIM_스크리닝_US_2026-10-09.xlsx"


class Test미국판정:
    def test_야후_재무로_1차(self, monkeypatch, quarterly, annual):
        monkeypatch.setattr(screener.yahoo, "load_financials", lambda code: (quarterly, annual))
        r = screener.screen_ca(UniverseItem("NVDA", "NVIDIA", "US", 4e12))
        assert r.status == PASS and "$" in r.detail     # 금액이 달러로 표시된다

    def test_I만_판단불가면_조건부_충족(self):
        from src.models import CanslimItem, CanslimResult, Verdict
        items = [CanslimItem(L, L, "", "", Verdict.PASS) for L in "CANSLM"]
        items.insert(5, CanslimItem("I", "I", "", "", Verdict.UNKNOWN))
        assert screener._qualified_ex_i(CanslimResult(items=items))
        items[0] = CanslimItem("C", "C", "", "", Verdict.FAIL)
        assert not screener._qualified_ex_i(CanslimResult(items=items))

    def test_엑셀은_달러_단위와_조건부_충족(self, tmp_path, monkeypatch, quarterly, annual):
        monkeypatch.setattr(screener.yahoo, "load_financials", lambda code: (quarterly, annual))
        monkeypatch.setattr(screener, "US_INTERVAL", 0)
        monkeypatch.setattr(screener, "deep_check", lambda item: screener.DeepResult(
            item, summary="미충족 (자료 부족)", qualified=False, qualified_ex_i=True))
        result = screener.run(universe=[UniverseItem("NVDA", "NVIDIA", "US", 4.2e12)],
                              progress=lambda m: None)
        wb = load_workbook(write_excel(result, tmp_path / "us.xlsx", currency="USD"))
        head = [c.value for c in wb["C·A 통과"][1]]
        assert "시가총액(백만$)" in head
        assert wb["C·A 통과"].cell(row=2, column=head.index("시가총액(백만$)") + 1).value == 4_200_000
        assert [r[1] for r in wb["CANSLIM 충족"].iter_rows(min_row=2, values_only=True)] == ["NVIDIA"]
        summary = {r[0]: r[1] for r in wb["요약"].iter_rows(values_only=True)}
        assert "I 기관 수급을 뺀" in summary["2차 CANSLIM 충족"]


class Test야후재무저장:
    def test_저장본이_있으면_야후를_부르지_않는다(self, tmp_path, monkeypatch, quarterly, annual):
        import json
        from src import yahoo
        path = tmp_path / "NVDA.json"
        path.write_text(json.dumps({"q": yahoo._table_to_dict(quarterly), "a": yahoo._table_to_dict(annual)}),
                        encoding="utf-8")
        import yfinance
        monkeypatch.setattr(yfinance, "Ticker", lambda s: (_ for _ in ()).throw(AssertionError("호출됨")))
        q, a = yahoo.load_financials("NVDA", cache_dir=tmp_path)
        assert [p.key for p in q.periods] == [p.key for p in quarterly.periods]
        assert q.rows == quarterly.rows and a.money_unit == annual.money_unit


class Test미국종목명:
    def test_꼬리말을_정리한다(self):
        assert screener.clean_us_name("Apple Inc. Common Stock") == "Apple Inc."
        assert screener.clean_us_name("Alphabet Inc. Class A Common Stock") == "Alphabet Inc. Class A"
        assert screener.clean_us_name(
            "Arm Holdings plc American Depositary Shares (each representing 1 share)") == "Arm Holdings plc (ADR)"
        assert screener.clean_us_name("Berkshire Hathaway Inc.") == "Berkshire Hathaway Inc."
