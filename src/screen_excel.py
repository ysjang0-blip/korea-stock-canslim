"""스크리너 결과를 엑셀로. 시트: 요약 / CANSLIM 충족 / C·A 통과 / C·A 근접."""

from __future__ import annotations

import datetime as dt
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import canslim
from .screener import ERROR, FAIL, NEAR, PASS, CaResult, DeepResult, ScreenResult

LETTERS = list("CANSLIM")
VERDICT_FILL = {
    "합격": PatternFill("solid", fgColor="D7F2DF"),
    "불합격": PatternFill("solid", fgColor="F8D7D7"),
    "판단불가": PatternFill("solid", fgColor="E8E8E8"),
}
HEADER_FILL = PatternFill("solid", fgColor="2F3B4C")
HEADER_FONT = Font(bold=True, color="FFFFFF")
PCT = '+#,##0.0"%";-#,##0.0"%"'


def _pct(v: float | None):
    return round(v, 1) if v is not None else None


CAP_UNIT = {"KRW": ("시가총액(억원)", 1e8), "USD": ("시가총액(백만$)", 1e6)}


def _base_cols(r: CaResult) -> list:
    i = r.item
    scale = CAP_UNIT[i.currency][1]
    return [i.market, i.name, i.code, round(i.cap / scale) if i.cap else None, r.latest_quarter,
            _pct(r.eps_yoy), _pct(r.revenue_yoy), _pct(r.annual_cagr), _pct(r.roe)]


def base_head(currency: str) -> list[str]:
    return ["시장", "종목명", "코드", CAP_UNIT[currency][0], "최근 분기",
            "C1 분기 EPS 전년비", "C4 분기 매출 전년비", "A1 연간 EPS 성장률", "A3 ROE"]


BASE_HEAD = base_head("KRW")
PCT_COLS = {6, 7, 8, 9}   # 1부터 센 열 번호 (BASE_HEAD 기준)


def _deep_cols(d: DeepResult | None) -> list:
    if d is None:
        return ["(2차 미실행)", ""] + [""] * 7 + ["", None]
    if d.error:
        return [f"오류: {d.error}", ""] + [""] * 7 + ["", None]
    return [d.summary, d.tally] + [d.verdicts.get(L, "") for L in LETTERS] + [d.opinion, _pct(d.base_gap)]


DEEP_HEAD = ["2차 종합", "합격·불합격·판단불가"] + LETTERS + ["투자의견(규칙)", "Base 목표가 괴리율"]


def _sheet(wb: Workbook, title: str, head: list[str], rows: list[list], widths: dict[str, int],
           pct_cols: set[int], verdict_cols: set[int]):
    ws = wb.create_sheet(title)
    ws.append(head)
    for row in rows:
        ws.append(row)
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "C2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(head))}{len(rows) + 1}"
    for idx, name in enumerate(head, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = widths.get(name, 11)
    for r in ws.iter_rows(min_row=2):
        for idx, cell in enumerate(r, start=1):
            if idx in pct_cols:
                cell.number_format = PCT
            if idx in verdict_cols and cell.value in VERDICT_FILL:
                cell.fill = VERDICT_FILL[cell.value]
                cell.alignment = Alignment(horizontal="center")
            if head[idx - 1] == "세부 조건":
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            if head[idx - 1].startswith("시가총액"):
                cell.number_format = "#,##0"
    ws.row_dimensions[1].height = 32
    return ws


WIDTHS = {"종목명": 22, "코드": 8, "시장": 8, "시가총액(억원)": 13, "시가총액(백만$)": 13, "최근 분기": 9,
          "2차 종합": 18, "합격·불합격·판단불가": 22, "투자의견(규칙)": 16,
          "세부 조건": 70, "모자란 조건": 16, **{L: 6 for L in LETTERS}}


def _sort_key(r: CaResult):
    return -(r.eps_yoy if r.eps_yoy is not None else -1e9)


def write_excel(result: ScreenResult, path: Path, currency: str = "KRW") -> Path:
    """currency='USD' 면 미국 결과 — 시가총액 단위와 '조건부 충족(I 제외)' 기준이 달라진다."""
    us = currency == "USD"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "요약"

    passed = sorted(result.by_status(PASS), key=_sort_key)
    near = sorted(result.by_status(NEAR), key=_sort_key)
    deep_head_start = len(BASE_HEAD) + 1
    verdict_cols = set(range(deep_head_start + 2, deep_head_start + 2 + len(LETTERS)))
    pct_all = PCT_COLS | {deep_head_start + 2 + len(LETTERS) + 1}

    def is_qualified(r: CaResult) -> bool:
        d = result.deep.get(r.item.code)
        return bool(d) and (d.qualified or (us and d.qualified_ex_i))

    qualified = [r for r in passed if is_qualified(r)]
    head = base_head(currency) + DEEP_HEAD + ["세부 조건"]

    def full_row(r: CaResult) -> list:
        return _base_cols(r) + _deep_cols(result.deep.get(r.item.code)) + [r.detail]

    _sheet(wb, "CANSLIM 충족", head, [full_row(r) for r in qualified], WIDTHS, pct_all, verdict_cols)
    _sheet(wb, "C·A 통과", head, [full_row(r) for r in passed], WIDTHS, pct_all, verdict_cols)
    _sheet(wb, "C·A 근접", base_head(currency) + ["모자란 조건", "세부 조건"],
           [_base_cols(r) + [", ".join([f"{c} 불합격" for c in r.failed] + [f"{c} 판단불가" for c in r.unknown]),
                             r.detail] for r in near],
           WIDTHS, PCT_COLS, set())

    started = dt.datetime.fromtimestamp(result.started)
    minutes = (result.finished - result.started) / 60 if result.finished else 0
    summary = [
        ("실행 시각", f"{started:%Y-%m-%d %H:%M} (소요 {minutes:.0f}분)"
         + (" — 중간에 중단됨, 일부 결과" if result.interrupted else "")),
        ("대상 종목", f"{result.universe:,}개 (미국 상장 보통주·ADR — 나스닥 스크리너 목록, 우선주·워런트 제외)"
                     if us else f"{result.universe:,}개 (코스피·코스닥 보통주 — ETF·ETN·우선주·스팩 제외)"),
        ("1차 C·A 통과", f"{len(passed):,}개"),
        ("1차 C·A 근접", f"{len(near):,}개 (조건 1개 불합격, 또는 자료가 없어 판단불가)"),
        ("1차 탈락", f"{len(result.by_status(FAIL)):,}개"),
        ("자료 오류", f"{len(result.by_status(ERROR)):,}개"),
        ("2차 CANSLIM 충족", f"{len(qualified):,}개 (I 기관 수급을 뺀 6개 항목 합격 — 미국은 I를 판정할 자료가 없음)"
                            if us else f"{len(qualified):,}개 (7개 항목 전부 합격)"),
        ("", ""),
        ("1차 기준 (웹앱과 같음)",
         f"C1 최근 분기 EPS 전년비 +{canslim.C_EPS_YOY_MIN:.0f}% 이상 · C3 직전 분기보다 EPS 증가 · "
         f"C4 분기 매출 전년비 +{canslim.C_REVENUE_YOY_MIN:.0f}% 이상 · A1 연간 EPS 연평균 +{canslim.A_CAGR_MIN:.0f}% 이상 · "
         f"A2 매년 EPS 증가 · A3 ROE {canslim.A_ROE_MIN:.0f}% 이상"),
        ("C2는 왜 1차에서 안 보나", "네이버 재무표가 확정 5분기뿐이라 직전 분기의 전년 동기가 없습니다. "
                               "1차는 판단불가를 허용하고, 2차에서 야후 실적 발표 이력으로 보강해 판정합니다."),
        ("2차", "1차 통과 종목만 웹앱과 같은 방식으로 7개 항목 전부 + 규칙 기반 투자의견·목표가를 계산"),
        ("정렬", "C1 분기 EPS 전년비가 큰 순서 (전년 동기 적자 → 흑자 전환은 증가율이 없어 아래쪽)"),
        ("주의", "기계적 필터입니다. 기저효과(전년 실적이 매우 작음)·일회성 이익을 걸러내지 못하니 "
               "세부 조건과 웹앱 상세 분석으로 확인하세요. 투자 판단 참고 자료이며 매수·매도 신호가 아닙니다."),
    ]
    for k, v in summary:
        ws.append([k, v])
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 110
    for row in ws.iter_rows():
        row[0].font = Font(bold=True)
        row[1].alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(path)
    return path
